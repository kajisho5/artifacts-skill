"""Regenerates the fixture corpus under tests/fixtures/<format>/.

Run with: python tests/fixtures/generate_fixtures.py
Requires the `dev` extra (`pip install -e ".[dev]"`) for `reportlab` and
`python-pptx`.

Per spec #31, the corpus deliberately includes broken files, not just
valid ones — the point is to test the *detection* of real problems.
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import pypdf
from docx import Document
from openpyxl import Workbook
from PIL import Image
from pptx import Presentation
from pypdf.generic import NameObject
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

# A handful of common install locations for a real TTF, used only to build
# the font-embedding fixtures below. Not vendored into the repo (keeps it
# free of binary font assets); if none of these exist, those two fixtures
# are skipped with a warning rather than failing the whole script — every
# other fixture is independent of having a system font available.
_CANDIDATE_TTF_PATHS = [
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf",
    "/System/Library/Fonts/Supplemental/Arial.ttf",
    "C:\\Windows\\Fonts\\arial.ttf",
]

PDF_OUT_DIR = Path(__file__).parent / "pdf"
PPTX_OUT_DIR = Path(__file__).parent / "pptx"
DOCX_OUT_DIR = Path(__file__).parent / "docx"
XLSX_OUT_DIR = Path(__file__).parent / "xlsx"
IMAGE_OUT_DIR = Path(__file__).parent / "image"
HTML_OUT_DIR = Path(__file__).parent / "html"
SVG_OUT_DIR = Path(__file__).parent / "svg"
CSV_OUT_DIR = Path(__file__).parent / "csv"
MARKDOWN_OUT_DIR = Path(__file__).parent / "markdown"
EPUB_OUT_DIR = Path(__file__).parent / "epub"
MEDIA_OUT_DIR = Path(__file__).parent / "media"


# ---------------------------------------------------------------- PDF ----

def make_good_2page() -> None:
    path = PDF_OUT_DIR / "good_2page.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setTitle("Sample Doc")
    c.drawString(100, 700, "Page one content")
    c.showPage()
    c.drawString(100, 700, "Page two content")
    c.showPage()
    c.save()


def make_blank_page_pdf() -> None:
    """3 pages, the middle one genuinely blank (no drawing calls at all) -
    a different failure mode from empty_0page.pdf (zero pages total).
    blank_pages must flag page 2 specifically."""
    path = PDF_OUT_DIR / "blank_page.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(100, 700, "Page one has content")
    c.showPage()
    c.showPage()  # page two: nothing drawn - genuinely blank
    c.drawString(100, 700, "Page three has content")
    c.showPage()
    c.save()


def make_leftover_placeholder_pdf() -> None:
    """A generated-looking PDF that still has unreviewed placeholder text -
    leftover_placeholder_text must WARN (or FAIL under
    forbid_placeholder_text), the PDF equivalent of the same fixture
    already built for DOCX/PPTX/HTML."""
    path = PDF_OUT_DIR / "leftover_placeholder.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.drawString(100, 700, "Click to add title")
    c.drawString(100, 680, "Lorem ipsum dolor sit amet, consectetur adipiscing elit.")
    c.drawString(100, 660, "TODO: write the real conclusion here.")
    c.showPage()
    c.save()


def make_empty_0page() -> None:
    path = PDF_OUT_DIR / "empty_0page.pdf"
    writer = pypdf.PdfWriter()
    with open(path, "wb") as f:
        writer.write(f)


def make_encrypted() -> None:
    src = PDF_OUT_DIR / "good_2page.pdf"
    path = PDF_OUT_DIR / "encrypted.pdf"
    reader = pypdf.PdfReader(str(src))
    writer = pypdf.PdfWriter()
    writer.append(reader)
    writer.encrypt("secret123")
    with open(path, "wb") as f:
        writer.write(f)


def make_corrupt_pdf() -> None:
    path = PDF_OUT_DIR / "corrupt.pdf"
    path.write_bytes(b"%PDF-1.4\n%% this is not a real pdf body, just garbage\nendobj trailer garbage")


def _find_system_ttf() -> str | None:
    for candidate in _CANDIDATE_TTF_PATHS:
        if Path(candidate).is_file():
            return candidate
    return None


def make_embedded_font_pdf() -> bool:
    """A page using a genuinely embedded TrueType font (not one of the
    standard 14) — verify_structural()'s font_embedding check should PASS."""
    ttf_path = _find_system_ttf()
    if ttf_path is None:
        return False
    path = PDF_OUT_DIR / "embedded_font.pdf"
    pdfmetrics.registerFont(TTFont("EmbeddedTestFont", ttf_path))
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("EmbeddedTestFont", 14)
    c.drawString(100, 700, "Embedded font test")
    c.showPage()
    c.save()
    return True


def make_nonembedded_custom_font_pdf() -> bool:
    """Same starting point as embedded_font.pdf, but with the embedded
    font's /FontFile2 stripped and its name changed to a plausible
    non-standard font — verify_structural()'s font_embedding check should
    WARN (or FAIL, with policy `forbid_unembedded_fonts`)."""
    embedded_path = PDF_OUT_DIR / "embedded_font.pdf"
    if not embedded_path.is_file():
        return False
    path = PDF_OUT_DIR / "nonembedded_custom_font.pdf"
    reader = pypdf.PdfReader(str(embedded_path))
    writer = pypdf.PdfWriter()
    writer.append(reader)
    page = writer.pages[0]
    for font_ref in page["/Resources"]["/Font"].values():
        font = font_ref.get_object()
        base_font = str(font.get("/BaseFont", ""))
        if "+" in base_font:  # the subset-tagged (embedded) font, not standard Helvetica
            font[NameObject("/BaseFont")] = NameObject("/CustomNonEmbedded")
            descriptor = font["/FontDescriptor"].get_object()
            for key in ("/FontFile", "/FontFile2", "/FontFile3"):
                if key in descriptor:
                    del descriptor[key]
            descriptor[NameObject("/FontName")] = NameObject("/CustomNonEmbedded")
    with open(path, "wb") as f:
        writer.write(f)
    return True


def make_mislabeled_html() -> None:
    path = PDF_OUT_DIR / "mislabeled_html.pdf"
    path.write_text("<html><body>not a pdf</body></html>\n")


# --------------------------------------------------------------- PPTX ----

def make_good_2slide_pptx() -> None:
    path = PPTX_OUT_DIR / "good_2slide.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = "Sample Deck"
    slide.placeholders[1].text = "Some body content"
    prs.slides.add_slide(prs.slide_layouts[6])  # blank layout, no placeholders
    prs.save(str(path))


def make_empty_placeholder_pptx() -> None:
    """One slide with a title placeholder left empty — the 'leftover
    placeholder' case verify_structural()'s empty_placeholders check
    should flag as WARN."""
    path = PPTX_OUT_DIR / "empty_placeholder.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.placeholders[1].text = "Body has content but the title was never filled in"
    prs.save(str(path))


def make_leftover_placeholder_text_pptx() -> None:
    """Placeholders are filled in (not empty - a different, already-covered
    case), but with unreviewed generation leftovers -
    leftover_placeholder_text must WARN."""
    path = PPTX_OUT_DIR / "leftover_placeholder_text.pptx"
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Click to add title"
    slide.placeholders[1].text = "Lorem ipsum dolor sit amet."
    prs.save(str(path))


def make_zero_slide_pptx() -> None:
    path = PPTX_OUT_DIR / "zero_slide.pptx"
    prs = Presentation()
    prs.save(str(path))


def make_corrupt_pptx() -> None:
    path = PPTX_OUT_DIR / "corrupt.pptx"
    path.write_bytes(b"PK\x03\x04this is not a real zip/ooxml body, just garbage bytes")


def make_mislabeled_pdf_as_pptx() -> None:
    """A real PDF saved with a .pptx extension, mirroring the PDF corpus's
    mislabeled_html.pdf — proves type detection never trusts the extension."""
    path = PPTX_OUT_DIR / "mislabeled_pdf.pptx"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# --------------------------------------------------------------- DOCX ----

def make_good_docx() -> None:
    path = DOCX_OUT_DIR / "good.docx"
    doc = Document()
    doc.add_heading("Sample Document", level=1)
    doc.add_paragraph("Some body content.")
    doc.add_table(rows=1, cols=2)
    doc.save(str(path))


def make_leftover_placeholder_docx() -> None:
    """A generated-looking document that still has unreviewed placeholder
    text — leftover_placeholder_text must WARN (or FAIL under
    forbid_placeholder_text)."""
    path = DOCX_OUT_DIR / "leftover_placeholder.docx"
    doc = Document()
    doc.add_heading("Click to add title", level=1)
    doc.add_paragraph("Lorem ipsum dolor sit amet, consectetur adipiscing elit.")
    doc.add_paragraph("TODO: write the real conclusion here.")
    doc.save(str(path))


def make_empty_docx() -> None:
    """Zero paragraphs — python-docx always creates at least the body
    element, but a document with no add_paragraph()/add_heading() calls
    has zero *paragraph* entries, which is the case verify_structural()'s
    paragraph_count check should flag as FAIL."""
    path = DOCX_OUT_DIR / "empty.docx"
    doc = Document()
    # Remove the single implicit empty paragraph python-docx starts with.
    for p in list(doc.paragraphs):
        p._element.getparent().remove(p._element)
    doc.save(str(path))


def make_corrupt_docx() -> None:
    path = DOCX_OUT_DIR / "corrupt.docx"
    path.write_bytes(b"PK\x03\x04this is not a real zip/ooxml body, just garbage bytes")


def make_mislabeled_pdf_as_docx() -> None:
    path = DOCX_OUT_DIR / "mislabeled_pdf.docx"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# --------------------------------------------------------------- XLSX ----

def _rewrite_zip_member(path: Path, member: str, old: str, new: str) -> None:
    """openpyxl can't write a cached formula result directly (it only ever
    writes the formula text) — so build a normal workbook, then patch the
    one cell's raw XML to add the `<v>` a real spreadsheet app would have
    written. This is the only way to get a deterministic "cached value
    present" or "cached error present" fixture without depending on
    LibreOffice (whose reliability in this exact dev environment is
    exactly what adapters/pptx/adapter.py's docstring documents as flaky)."""
    with zipfile.ZipFile(path, "r") as zf:
        items = {name: zf.read(name) for name in zf.namelist()}
    xml = items[member].decode("utf-8")
    assert old in xml, f"expected substring not found in {member}: {old!r}"
    items[member] = xml.replace(old, new, 1).encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)


def _make_formula_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Label"
    ws["B1"] = 10
    ws["B2"] = 20
    ws["B3"] = "=B1+B2"
    wb.create_sheet("Sheet2")
    wb.save(str(path))


def make_good_xlsx() -> None:
    """Two sheets, one formula, a real (non-error) cached result — the
    PASS branch of formula_cached_errors, and the ordinary UNKNOWN branch
    of formula_recalculation (openpyxl never recalculates, regardless of
    what's cached)."""
    path = XLSX_OUT_DIR / "good.xlsx"
    _make_formula_workbook(path)
    _rewrite_zip_member(path, "xl/worksheets/sheet1.xml", '<c r="B3"><f>B1+B2</f><v></v></c>', '<c r="B3"><f>B1+B2</f><v>30</v></c>')


def make_formula_error_xlsx() -> None:
    """A formula cell with a cached #REF! result — must FAIL formula_cached_errors."""
    path = XLSX_OUT_DIR / "formula_error.xlsx"
    _make_formula_workbook(path)
    _rewrite_zip_member(
        path, "xl/worksheets/sheet1.xml",
        '<c r="B3"><f>B1+B2</f><v></v></c>', '<c r="B3" t="e"><f>B1+B2</f><v>#REF!</v></c>',
    )


def make_no_formula_xlsx() -> None:
    """No formulas at all — formula_cached_errors/formula_recalculation
    should both be SKIPPED, not UNKNOWN (a genuinely not-applicable case
    is a different, more benign signal than "we couldn't check")."""
    path = XLSX_OUT_DIR / "no_formula.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Data"
    ws["A1"] = "Label"
    ws["B1"] = 10
    wb.save(str(path))


def make_entity_bomb_xlsx() -> None:
    """A DOCTYPE declaring a custom entity, injected into the worksheet
    XML part — not an actual expansion bomb (a real one would be
    unpleasant to keep in a test fixture directory), just enough to prove
    security/xml_safety.py's reject_xml_entities_in_zip() catches the
    pattern before openpyxl ever opens the file (Issue #21: confirmed by
    direct testing that openpyxl's worksheet reader resolves and amplifies
    exactly this shape when nothing intercepts it first)."""
    path = XLSX_OUT_DIR / "entity_bomb.xlsx"
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    wb.save(str(path))
    with zipfile.ZipFile(path, "r") as zf:
        items = {name: zf.read(name) for name in zf.namelist()}
    sheet_xml = items["xl/worksheets/sheet1.xml"].decode("utf-8")
    injected = sheet_xml.replace(
        "<worksheet ",
        '<!DOCTYPE worksheet [<!ENTITY xxe "PWNED">]>\n<worksheet ',
        1,
    )
    items["xl/worksheets/sheet1.xml"] = injected.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)


def make_leftover_placeholder_xlsx() -> None:
    """A cell still holds unreviewed placeholder text - leftover_placeholder_text
    must WARN (or FAIL under forbid_placeholder_text), the XLSX equivalent
    of the same fixture already built for DOCX/PPTX/HTML/PDF."""
    path = XLSX_OUT_DIR / "leftover_placeholder.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Click to add title"
    ws["A2"] = "TODO: fill in real Q3 numbers"
    ws["A3"] = "Lorem ipsum dolor sit amet"
    wb.save(str(path))


def make_external_link_xlsx() -> None:
    """A workbook with a real `<externalReferences>` part pointing at
    another workbook file — must produce a non-empty `external_links` in
    inspect() and WARN/FAIL the `external_links` structural check.

    openpyxl has no write-side API for this (confirmed: saving a formula
    string like "=[1]Sheet1!A1" does NOT create the externalLinks parts —
    it's just treated as opaque formula text), so this hand-builds the
    three OOXML parts a real external reference needs and wires them into
    the zip the same way _rewrite_zip_member patches a single cell:
    xl/externalLinks/externalLink1.xml (the external book + sheet names),
    its _rels file (the actual external target path), plus registering
    both in xl/workbook.xml's <externalReferences>, xl/_rels/workbook.xml.rels,
    and [Content_Types].xml. Verified by round-tripping through openpyxl's
    own reader (wb._external_links is populated) before trusting this as a
    fixture, not assumed correct from the XML alone.
    """
    path = XLSX_OUT_DIR / "external_link.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Label"
    ws["B1"] = "=[1]Sheet1!A1"  # references the (not-actually-present) external book
    wb.save(str(path))

    with zipfile.ZipFile(path, "r") as zf:
        items = {name: zf.read(name) for name in zf.namelist()}

    wb_xml = items["xl/workbook.xml"].decode()
    assert "</sheets>" in wb_xml
    items["xl/workbook.xml"] = wb_xml.replace(
        "</sheets>",
        '</sheets><externalReferences><externalReference '
        'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId4"/>'
        "</externalReferences>",
        1,
    ).encode()

    rels_xml = items["xl/_rels/workbook.xml.rels"].decode()
    assert "</Relationships>" in rels_xml
    items["xl/_rels/workbook.xml.rels"] = rels_xml.replace(
        "</Relationships>",
        '<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink" '
        'Target="externalLinks/externalLink1.xml" Id="rId4"/></Relationships>',
    ).encode()

    ct_xml = items["[Content_Types].xml"].decode()
    assert "</Types>" in ct_xml
    items["[Content_Types].xml"] = ct_xml.replace(
        "</Types>",
        '<Override PartName="/xl/externalLinks/externalLink1.xml" '
        'ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml"/></Types>',
    ).encode()

    items["xl/externalLinks/externalLink1.xml"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<externalLink xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">'
        b'<externalBook xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" r:id="rId1">'
        b"<sheetNames><sheetName val=\"Sheet1\"/></sheetNames>"
        b"</externalBook></externalLink>"
    )
    items["xl/externalLinks/_rels/externalLink1.xml.rels"] = (
        b'<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId1" '
        b'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
        b'Target="other_workbook.xlsx" TargetMode="External"/></Relationships>'
    )

    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items.items():
            zf.writestr(name, data)

    # Don't trust the hand-built XML blindly - confirm openpyxl's own
    # reader actually populates _external_links from it before shipping
    # this as a fixture other tests will rely on.
    from openpyxl import load_workbook

    reloaded = load_workbook(str(path))
    assert getattr(reloaded, "_external_links", None), "external_link.xlsx fixture failed its own round-trip check"


def make_corrupt_xlsx() -> None:
    path = XLSX_OUT_DIR / "corrupt.xlsx"
    path.write_bytes(b"PK\x03\x04this is not a real zip/ooxml body, just garbage bytes")


def make_mislabeled_pdf_as_xlsx() -> None:
    path = XLSX_OUT_DIR / "mislabeled_pdf.xlsx"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# -------------------------------------------------------------- IMAGE ----

def make_good_png() -> None:
    path = IMAGE_OUT_DIR / "good.png"
    Image.new("RGB", (200, 100), color=(255, 0, 0)).save(path)


def make_alpha_png() -> None:
    path = IMAGE_OUT_DIR / "alpha.png"
    Image.new("RGBA", (150, 150), color=(0, 255, 0, 128)).save(path)


def make_exif_rotated_jpeg() -> None:
    """EXIF orientation 6 (rotate 90° CW) — pure Pillow, no extra
    dependency needed; Image.Exif() round-trips through JPEG save/load."""
    path = IMAGE_OUT_DIR / "exif_rotated.jpg"
    img = Image.new("RGB", (300, 200), color=(0, 0, 255))
    exif = img.getexif()
    exif[274] = 6
    img.save(path, exif=exif)


def make_corrupt_png() -> None:
    path = IMAGE_OUT_DIR / "corrupt.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\nthis is not a real PNG body, just garbage bytes")


def make_mislabeled_pdf_as_png() -> None:
    path = IMAGE_OUT_DIR / "mislabeled_pdf.png"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# --------------------------------------------------------------- HTML ----

def make_good_html() -> None:
    path = HTML_OUT_DIR / "good.html"
    path.write_text(
        "<!doctype html>\n<html><head><title>Sample Page</title></head>\n"
        "<body><h1>Hello Artifact Skill</h1><p>Some body content.</p></body></html>\n"
    )


def make_missing_local_resource_html() -> None:
    path = HTML_OUT_DIR / "missing_local_resource.html"
    path.write_text(
        "<!doctype html>\n<html><head><title>Broken Refs</title>"
        '<link rel="stylesheet" href="does-not-exist.css"></head>\n'
        '<body><img src="also-missing.png"></body></html>\n'
    )


def make_external_resource_html() -> None:
    path = HTML_OUT_DIR / "external_resource.html"
    path.write_text(
        "<!doctype html>\n<html><head><title>External Refs</title></head>\n"
        '<body><img src="https://example.com/some-image.png"></body></html>\n'
    )


def make_leftover_placeholder_html() -> None:
    path = HTML_OUT_DIR / "leftover_placeholder.html"
    path.write_text(
        "<!doctype html>\n<html><head><title>Sample Page</title></head>\n"
        "<body><h1>Click to add title</h1>"
        "<p>Lorem ipsum dolor sit amet, consectetur adipiscing elit.</p>"
        "<script>// TODO: this is JS code, not document text - must not be flagged</script>"
        "</body></html>\n"
    )


def make_no_title_html() -> None:
    path = HTML_OUT_DIR / "no_title.html"
    path.write_text("<!doctype html>\n<html><body><p>No title element here.</p></body></html>\n")


def make_binary_garbage_html() -> None:
    path = HTML_OUT_DIR / "binary_garbage.html"
    path.write_bytes(b"\xff\xfe\x00\x01not valid utf-8 \xfe\xff pretending to be html")


def make_mislabeled_pdf_as_html() -> None:
    path = HTML_OUT_DIR / "mislabeled_pdf.html"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# ---------------------------------------------------------------- SVG ----

def make_good_svg() -> None:
    path = SVG_OUT_DIR / "good.svg"
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">'
        '<circle cx="50" cy="50" r="40" fill="red"/></svg>\n'
    )


def make_leftover_placeholder_svg() -> None:
    """A <text> element still holds unreviewed placeholder text -
    leftover_placeholder_text must WARN (or FAIL under
    forbid_placeholder_text), the SVG equivalent of the same fixture
    already built for DOCX/PPTX/HTML/PDF/XLSX."""
    path = SVG_OUT_DIR / "leftover_placeholder.svg"
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="200" height="100">'
        '<text x="10" y="20">Click to add title</text>'
        '<text x="10" y="40">Lorem ipsum dolor sit amet</text>'
        '<text x="10" y="60">TODO: replace with real copy</text>'
        "</svg>\n"
    )


def make_missing_local_resource_svg() -> None:
    path = SVG_OUT_DIR / "missing_local_resource.svg"
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">'
        '<image href="does-not-exist.png" width="10" height="10"/></svg>\n'
    )


def make_external_resource_svg() -> None:
    path = SVG_OUT_DIR / "external_resource.svg"
    path.write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100">'
        '<image href="https://example.com/some-image.png" width="10" height="10"/></svg>\n'
    )


def make_no_size_svg() -> None:
    path = SVG_OUT_DIR / "no_size.svg"
    path.write_text('<svg xmlns="http://www.w3.org/2000/svg"><circle cx="5" cy="5" r="4"/></svg>\n')


def make_malformed_svg() -> None:
    path = SVG_OUT_DIR / "malformed.svg"
    path.write_text('<svg xmlns="http://www.w3.org/2000/svg"><circle cx="5" cy="5" r="4"></svg>\n')  # unclosed <circle>


def make_entity_bomb_svg() -> None:
    """A DOCTYPE declaring a custom entity — not an actual expansion bomb
    (a real one would be unpleasant to keep in a test fixture directory),
    just enough to prove _reject_xml_entities() catches the pattern before
    any parsing is attempted."""
    path = SVG_OUT_DIR / "entity_bomb.svg"
    path.write_text(
        '<?xml version="1.0"?>\n<!DOCTYPE svg [<!ENTITY lol "lol">]>\n'
        '<svg xmlns="http://www.w3.org/2000/svg"><title>&lol;</title></svg>\n'
    )


def make_mislabeled_pdf_as_svg() -> None:
    path = SVG_OUT_DIR / "mislabeled_pdf.svg"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# ---------------------------------------------------------------- CSV ----

def make_good_csv() -> None:
    path = CSV_OUT_DIR / "good.csv"
    path.write_text("name,age,city\nAlice,30,Tokyo\nBob,25,Osaka\nCarol,40,Kyoto\n")


def make_ragged_csv() -> None:
    """Row 2 (0-indexed) has one fewer field than the header - column_count_consistency must FAIL."""
    path = CSV_OUT_DIR / "ragged.csv"
    path.write_text("a,b,c\n1,2,3\n4,5\n6,7,8\n")


def make_leftover_placeholder_csv() -> None:
    path = CSV_OUT_DIR / "leftover_placeholder.csv"
    path.write_text("item,note\nWidget,TODO fill in real price\nGadget,Lorem ipsum dolor sit amet\n")


def make_mislabeled_pdf_as_csv() -> None:
    path = CSV_OUT_DIR / "mislabeled_pdf.csv"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# ----------------------------------------------------------- MARKDOWN ----

def make_good_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "good.md"
    path.write_text(
        "# Title\n\nSome intro text with a [link](good.md) to itself.\n\n"
        "- item one\n- item two\n\n```python\nprint('hello')\n```\n"
    )


def make_unclosed_fence_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "unclosed_fence.md"
    path.write_text("# Title\n\n```python\nprint('never closed')\n\nMore text after.\n")


def make_missing_local_resource_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "missing_local_resource.md"
    path.write_text("# Title\n\nSee ![diagram](does-not-exist.png) for details.\n")


def make_external_resource_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "external_resource.md"
    path.write_text("# Title\n\nSee [the spec](https://example.com/spec) for details.\n")


def make_leftover_placeholder_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "leftover_placeholder.md"
    path.write_text(
        "# Title\n\nTODO: write the real introduction.\n\n- Lorem ipsum dolor sit amet.\n- Another point.\n"
    )


def make_leftover_in_code_fence_markdown() -> None:
    """A "TODO" that appears only inside a fenced code block - the code
    fence's *body* must be excluded from the leftover-text scan, the same
    "code is not document text" reasoning the HTML adapter applies to
    <script>/<style> content."""
    path = MARKDOWN_OUT_DIR / "leftover_in_code_fence.md"
    path.write_text("# Title\n\nReal, reviewed content here.\n\n```text\n# TODO: this is sample output, not a note\n```\n")


def make_mislabeled_pdf_as_markdown() -> None:
    path = MARKDOWN_OUT_DIR / "mislabeled_pdf.md"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# --------------------------------------------------------------- EPUB ----

_EPUB_CONTAINER_XML = (
    '<?xml version="1.0"?>\n'
    '<container xmlns="urn:oasis:names:tc:opendocument:xmlns:container" version="1.0">\n'
    '  <rootfiles><rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/></rootfiles>\n'
    "</container>\n"
)


def _epub_opf(title: str, creator: str, chapter_href: str, spine_idref: str) -> str:
    return (
        '<?xml version="1.0"?>\n'
        '<package xmlns="http://www.idpf.org/2007/opf" version="3.0" unique-identifier="bookid">\n'
        '  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/">\n'
        f"    <dc:title>{title}</dc:title>\n"
        f"    <dc:creator>{creator}</dc:creator>\n"
        "    <dc:language>en</dc:language>\n"
        '    <dc:identifier id="bookid">urn:uuid:00000000-0000-0000-0000-000000000000</dc:identifier>\n'
        "  </metadata>\n"
        "  <manifest>\n"
        f'    <item id="chap1" href="{chapter_href}" media-type="application/xhtml+xml"/>\n'
        "  </manifest>\n"
        "  <spine>\n"
        f'    <itemref idref="{spine_idref}"/>\n'
        "  </spine>\n"
        "</package>\n"
    )


def _write_epub(path: Path, *, chapter_text: str, chapter_href: str = "text/chapter1.xhtml",
                 manifest_href: str | None = None, spine_idref: str = "chap1",
                 title: str = "My Book", creator: str = "Author Name",
                 mimetype_first_and_stored: bool = True, opf_override: str | None = None) -> None:
    with zipfile.ZipFile(path, "w") as zf:
        if mimetype_first_and_stored:
            zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", zipfile.ZIP_STORED)
        zf.writestr("META-INF/container.xml", _EPUB_CONTAINER_XML)
        zf.writestr(f"OEBPS/{chapter_href}", f"<html><body><p>{chapter_text}</p></body></html>")
        opf = opf_override or _epub_opf(title, creator, manifest_href or chapter_href, spine_idref)
        zf.writestr("OEBPS/content.opf", opf)
        if not mimetype_first_and_stored:
            # written last and compressed - fails the OCF "first, stored" requirement (WARN, not FAIL).
            zf.writestr("mimetype", "application/epub+zip")


def make_good_epub() -> None:
    _write_epub(EPUB_OUT_DIR / "good.epub", chapter_text="Hello world, this is chapter one.")


def make_broken_manifest_epub() -> None:
    _write_epub(
        EPUB_OUT_DIR / "broken_manifest.epub",
        chapter_text="Hello.", manifest_href="text/does-not-exist.xhtml",
    )


def make_broken_spine_epub() -> None:
    _write_epub(EPUB_OUT_DIR / "broken_spine.epub", chapter_text="Hello.", spine_idref="no-such-id")


def make_leftover_placeholder_epub() -> None:
    _write_epub(
        EPUB_OUT_DIR / "leftover_placeholder.epub",
        chapter_text="TODO: write the real chapter content. Lorem ipsum dolor sit amet.",
    )


def make_mimetype_not_first_epub() -> None:
    _write_epub(EPUB_OUT_DIR / "mimetype_not_first.epub", chapter_text="Hello.", mimetype_first_and_stored=False)


def make_entity_bomb_epub() -> None:
    """DOCTYPE declaring a custom entity in the OPF - proves the per-member
    reject_xml_entity_declaration() guard runs before ET.fromstring() on
    the package document, the same pattern as SVG's entity_bomb fixture."""
    opf = (
        '<?xml version="1.0"?>\n<!DOCTYPE package [<!ENTITY lol "lol">]>\n'
        '<package xmlns="http://www.idpf.org/2007/opf">\n'
        '  <metadata xmlns:dc="http://purl.org/dc/elements/1.1/"><dc:title>&lol;</dc:title></metadata>\n'
        "  <manifest/><spine/>\n</package>\n"
    )
    _write_epub(EPUB_OUT_DIR / "entity_bomb.epub", chapter_text="Hello.", opf_override=opf)


def make_missing_container_epub() -> None:
    """No META-INF/container.xml at all - a zip with the right mimetype
    member (so it's still *detected* as EPUB) but structurally unreadable."""
    path = EPUB_OUT_DIR / "missing_container.epub"
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", zipfile.ZIP_STORED)
        zf.writestr("OEBPS/content.opf", _epub_opf("Title", "Author", "text/chapter1.xhtml", "chap1"))


def make_mislabeled_pdf_as_epub() -> None:
    path = EPUB_OUT_DIR / "mislabeled_pdf.epub"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


# --------------------------------------------------------------- Media ----
# Unlike every format above, these are generated by shelling out to a real
# `ffmpeg` binary rather than a pure-Python library - there is no
# pip-installable way to synthesize a real MP4/WebM/WAV. Same graceful-skip
# shape as the embedded-font PDF fixtures above: if ffmpeg isn't on PATH,
# every make_*() below returns False and main() prints one warning instead
# of failing the whole script.


def _ffmpeg_binary() -> str | None:
    return shutil.which("ffmpeg")


def _run_ffmpeg(args: list[str]) -> bool:
    ffmpeg = _ffmpeg_binary()
    if ffmpeg is None:
        return False
    proc = subprocess.run(
        [ffmpeg, "-y", "-loglevel", "error", *args], capture_output=True, text=True, check=False,
    )
    return proc.returncode == 0


def make_good_mp4() -> bool:
    """2s, 320x240, H.264 video + AAC audio - the "everything present and
    healthy" case verify_structural() should PASS cleanly."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=10",
        "-f", "lavfi", "-i", "sine=frequency=1000:duration=2",
        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-c:a", "aac", "-shortest",
        str(MEDIA_OUT_DIR / "good.mp4"),
    ])


def make_good_webm() -> bool:
    """Video-only WebM (VP9, no audio track) - exercises has_audio=False
    on an otherwise healthy file, and the EBML magic-byte detection path."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "testsrc=duration=2:size=320x240:rate=10",
        "-c:v", "libvpx-vp9", "-pix_fmt", "yuv420p",
        str(MEDIA_OUT_DIR / "good.webm"),
    ])


def make_good_wav() -> bool:
    """Audio-only WAV (RIFF/WAVE magic bytes, no ftyp/EBML at all) -
    render() must report zero files with a warning, not an error, since
    there is no video stream to extract a frame from."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "sine=frequency=1000:duration=2",
        "-c:a", "pcm_s16le",
        str(MEDIA_OUT_DIR / "good.wav"),
    ])


def make_audio_only_mp4() -> bool:
    """An MP4/M4A-family container (ftyp magic bytes) that is genuinely
    audio-only - proves has_video/has_audio are read from the real probed
    streams, not assumed from the container family."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "sine=frequency=440:duration=1",
        "-c:a", "aac",
        str(MEDIA_OUT_DIR / "audio_only.mp4"),
    ])


def make_leftover_placeholder_wav() -> bool:
    """A container-metadata title tag containing an obvious leftover
    marker - exercises the same leftover-placeholder-text check every
    other adapter has, applied to ffprobe's format.tags instead of
    document body text."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "sine=frequency=1000:duration=1",
        "-metadata", "title=TODO fix this later",
        "-c:a", "pcm_s16le",
        str(MEDIA_OUT_DIR / "leftover_placeholder.wav"),
    ])


def make_audio_with_cover_m4a() -> bool:
    """Adversarial-review finding: an audio file with embedded cover art
    (extremely common - iTunes/Apple Music/podcast-tool M4A, ripped MP3/
    M4A with album art) gets an mjpeg "video" stream from ffprobe
    alongside the real audio stream, marked `disposition.attached_pic=1`.
    Without excluding it, this used to be misreported as has_video=True
    and crashed render() (no real frame to seek to) instead of taking the
    ordinary audio-only path - see adapters/media/adapter.py's
    _first_stream() for the fix this fixture regression-tests."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "sine=frequency=1000:duration=2",
        "-f", "lavfi", "-i", "color=c=blue:s=64x64:d=1",
        "-map", "0:a", "-map", "1:v", "-c:a", "aac", "-c:v", "mjpeg", "-disposition:v", "attached_pic",
        str(MEDIA_OUT_DIR / "audio_with_cover.m4a"),
    ])


def make_single_frame_mp4() -> bool:
    """Adversarial-review finding: a short/single-frame-ish clip can
    genuinely have no frame at render()'s computed midpoint seek time
    (only a real frame at t=0), so ffmpeg exits 0 with no output there
    even though the file is not corrupt - see adapters/media/adapter.py's
    render() for the retry-at-t=0 fallback this fixture regression-tests."""
    return _run_ffmpeg([
        "-f", "lavfi", "-i", "color=c=red:s=64x64:d=0.01:r=25",
        "-c:v", "libx264", "-pix_fmt", "yuv420p", "-frames:v", "1",
        str(MEDIA_OUT_DIR / "single_frame.mp4"),
    ])


def make_corrupt_truncated_mp4() -> bool:
    """A real MP4 truncated mid-stream - large enough to still look like a
    plausible file, but ffprobe should still refuse it as unreadable."""
    good = MEDIA_OUT_DIR / "good.mp4"
    if not good.is_file():
        return False
    (MEDIA_OUT_DIR / "corrupt_truncated.mp4").write_bytes(good.read_bytes()[:2000])
    return True


def make_garbage_too_short_mp4() -> bool:
    """Just the first few bytes of a real MP4 - still matches the 'ftyp'
    magic-byte check at the type-detection layer, but ffprobe has nowhere
    near enough data to read anything from it."""
    good = MEDIA_OUT_DIR / "good.mp4"
    if not good.is_file():
        return False
    (MEDIA_OUT_DIR / "garbage_too_short.mp4").write_bytes(good.read_bytes()[:20])
    return True


def make_mislabeled_pdf_as_mp4() -> None:
    path = MEDIA_OUT_DIR / "mislabeled_pdf.mp4"
    path.write_bytes((PDF_OUT_DIR / "good_2page.pdf").read_bytes())


if __name__ == "__main__":
    PDF_OUT_DIR.mkdir(parents=True, exist_ok=True)
    PPTX_OUT_DIR.mkdir(parents=True, exist_ok=True)
    DOCX_OUT_DIR.mkdir(parents=True, exist_ok=True)
    XLSX_OUT_DIR.mkdir(parents=True, exist_ok=True)
    IMAGE_OUT_DIR.mkdir(parents=True, exist_ok=True)
    HTML_OUT_DIR.mkdir(parents=True, exist_ok=True)
    SVG_OUT_DIR.mkdir(parents=True, exist_ok=True)
    CSV_OUT_DIR.mkdir(parents=True, exist_ok=True)
    MARKDOWN_OUT_DIR.mkdir(parents=True, exist_ok=True)
    EPUB_OUT_DIR.mkdir(parents=True, exist_ok=True)
    MEDIA_OUT_DIR.mkdir(parents=True, exist_ok=True)

    make_good_2page()
    make_blank_page_pdf()
    make_leftover_placeholder_pdf()
    make_empty_0page()
    make_encrypted()
    make_corrupt_pdf()
    make_mislabeled_html()
    if make_embedded_font_pdf():
        make_nonembedded_custom_font_pdf()
    else:
        print("WARNING: no system TTF found in _CANDIDATE_TTF_PATHS; skipped "
              "embedded_font.pdf / nonembedded_custom_font.pdf (leaving any existing copies as-is).")

    make_good_2slide_pptx()
    make_empty_placeholder_pptx()
    make_leftover_placeholder_text_pptx()
    make_zero_slide_pptx()
    make_corrupt_pptx()
    make_mislabeled_pdf_as_pptx()

    make_good_docx()
    make_empty_docx()
    make_leftover_placeholder_docx()
    make_corrupt_docx()
    make_mislabeled_pdf_as_docx()

    make_good_xlsx()
    make_formula_error_xlsx()
    make_no_formula_xlsx()
    make_leftover_placeholder_xlsx()
    make_external_link_xlsx()
    make_entity_bomb_xlsx()
    make_corrupt_xlsx()
    make_mislabeled_pdf_as_xlsx()

    make_good_png()
    make_alpha_png()
    make_exif_rotated_jpeg()
    make_corrupt_png()
    make_mislabeled_pdf_as_png()

    make_good_html()
    make_missing_local_resource_html()
    make_external_resource_html()
    make_no_title_html()
    make_leftover_placeholder_html()
    make_binary_garbage_html()
    make_mislabeled_pdf_as_html()

    make_good_svg()
    make_missing_local_resource_svg()
    make_external_resource_svg()
    make_leftover_placeholder_svg()
    make_no_size_svg()
    make_malformed_svg()
    make_entity_bomb_svg()
    make_mislabeled_pdf_as_svg()

    make_good_csv()
    make_ragged_csv()
    make_leftover_placeholder_csv()
    make_mislabeled_pdf_as_csv()

    make_good_markdown()
    make_unclosed_fence_markdown()
    make_missing_local_resource_markdown()
    make_external_resource_markdown()
    make_leftover_placeholder_markdown()
    make_leftover_in_code_fence_markdown()
    make_mislabeled_pdf_as_markdown()

    make_good_epub()
    make_broken_manifest_epub()
    make_broken_spine_epub()
    make_leftover_placeholder_epub()
    make_mimetype_not_first_epub()
    make_entity_bomb_epub()
    make_missing_container_epub()
    make_mislabeled_pdf_as_epub()

    if make_good_mp4():
        make_corrupt_truncated_mp4()
        make_garbage_too_short_mp4()
    else:
        print("WARNING: no ffmpeg found on PATH; skipped good.mp4 and the fixtures derived from it "
              "(corrupt_truncated.mp4, garbage_too_short.mp4) — leaving any existing copies as-is.")
    make_good_webm()
    make_good_wav()
    make_audio_only_mp4()
    make_leftover_placeholder_wav()
    make_audio_with_cover_m4a()
    make_single_frame_mp4()
    make_mislabeled_pdf_as_mp4()

    print(f"Wrote PDF fixtures to {PDF_OUT_DIR}")
    print(f"Wrote PPTX fixtures to {PPTX_OUT_DIR}")
    print(f"Wrote DOCX fixtures to {DOCX_OUT_DIR}")
    print(f"Wrote XLSX fixtures to {XLSX_OUT_DIR}")
    print(f"Wrote image fixtures to {IMAGE_OUT_DIR}")
    print(f"Wrote HTML fixtures to {HTML_OUT_DIR}")
    print(f"Wrote SVG fixtures to {SVG_OUT_DIR}")
    print(f"Wrote CSV fixtures to {CSV_OUT_DIR}")
    print(f"Wrote Markdown fixtures to {MARKDOWN_OUT_DIR}")
    print(f"Wrote EPUB fixtures to {EPUB_OUT_DIR}")
    print(f"Wrote media fixtures to {MEDIA_OUT_DIR}")
