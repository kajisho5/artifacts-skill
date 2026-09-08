"""XLSX adapter.

Structural read/write is pure Python via `openpyxl` (MIT), no external
binary. Rendering converts XLSX -> PDF with LibreOffice headless via the
same shared `rendering/office_convert.py` the PPTX/DOCX adapters use.

## The recalculation decision (spec §14 / Issue #5)

`openpyxl` does not evaluate formulas. It can only report:
  - the formula text itself (`data_type == "f"`, e.g. `"=B1+B2"`), or
  - a *cached* result, if one was already stored in the file by whatever
    application (Excel, LibreOffice, ...) last saved it — and `None` if
    nothing ever computed one (e.g. a file `openpyxl` itself wrote, as its
    own test fixtures for this adapter demonstrate).

So there is no way for this adapter to answer "are these formulas
currently correct" from the XML alone. Rather than silently trusting a
possibly-stale cached value, or pretending recalculation happened, this
adapter reports two separate, honest facts (see `verify_structural()`):
  - `formula_cached_errors`: PASS/FAIL based on whatever cached results
    *do* exist (catches an error a source application already flagged,
    e.g. `#REF!`), `SKIPPED` if there are no formulas at all, `UNKNOWN` if
    there are formulas but no cached values to inspect.
  - `formula_recalculation`: `UNKNOWN` whenever any formula is present,
    unconditionally — this project does not attempt LibreOffice-macro-
    based recalculation (a real option, deliberately not taken for this
    MVP: it would mean shelling out to a macro-scripting interface with a
    much larger attack surface and failure-mode space than the plain
    `--convert-to pdf` LibreOffice usage the render path already needs,
    for a benefit — "are formulas fresh" — that is speculative until a
    real user asks for it). See docs/adapters.md and docs/roadmap.md
    Issue #5 for the fuller writeup of this trade-off.
"""

from __future__ import annotations

import importlib.util
import tempfile
import zipfile
from pathlib import Path
from typing import Any

from artifact_skill.adapters.base import ArtifactAdapter, OperationSpec, RenderResult
from artifact_skill.core.artifact import ArtifactRef, ArtifactType, InspectionReport
from artifact_skill.core.capability import Capability, CapabilityStatus
from artifact_skill.core.errors import ArtifactCapabilityError, ArtifactInputError
from artifact_skill.core.operation import OperationPlan
from artifact_skill.core.verification import Check, CheckStatus, VerificationResult
from artifact_skill.leftover_text import find_leftover_markers
from artifact_skill.rendering.office_convert import convert_to_pdf, soffice_binary
from artifact_skill.rendering.pdf_pages import render_pdf_pages
from artifact_skill.security.limits import DEFAULT_LIMITS, Limits
from artifact_skill.security.paths import atomic_copy, check_input_size
from artifact_skill.security.xml_safety import reject_xml_entities_in_zip

_ERROR_TOKENS = {"#REF!", "#VALUE!", "#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#SPILL!", "#CALC!"}


def _has(module: str) -> bool:
    return importlib.util.find_spec(module) is not None


def _require_openpyxl():
    if not _has("openpyxl"):
        raise ArtifactCapabilityError(
            code="ARTIFACT_CAPABILITY_MISSING",
            message="openpyxl is not installed; XLSX structural read/write is unavailable.",
            remediation="Install with: pip install 'artifacts-skill[xlsx]' (or `pip install openpyxl`).",
            evidence={"capability_id": "xlsx.structural"},
        )
    import openpyxl

    return openpyxl


def _reject_entities_before_opening(path: Path) -> None:
    """Scan the zip for a DOCTYPE/ENTITY declaration (Issue #21) before
    handing the file to openpyxl. Needed because openpyxl only uses its
    hardened lxml/defusedxml XML parser when one of those packages happens
    to be importable — this project's own `xlsx` extra pulls in neither,
    so a `pip install -e ".[xlsx]"`-only install gets zero protection from
    openpyxl itself; see `security/xml_safety.py`'s module docstring for
    the full audit this closes. A malformed (non-zip) file is left for the
    caller's own `openpyxl.load_workbook()` try/except to report as
    ARTIFACT_XLSX_UNREADABLE, the same as any other corruption.
    """
    try:
        reject_xml_entities_in_zip(path)
    except zipfile.BadZipFile:
        pass


class XlsxAdapter(ArtifactAdapter):
    id = "xlsx"
    artifact_type = ArtifactType.XLSX

    @classmethod
    def detect(cls, ref: ArtifactRef) -> bool:
        return ref.type == ArtifactType.XLSX

    def operations(self) -> dict[str, OperationSpec]:
        return {
            "metadata_set": OperationSpec(
                name="metadata_set",
                description="Set core workbook properties (title, author, subject, keywords).",
                args_schema={
                    "type": "object",
                    "properties": {
                        "title": {"type": "string"},
                        "author": {"type": "string"},
                        "subject": {"type": "string"},
                        "keywords": {"type": "string"},
                    },
                    "additionalProperties": False,
                },
                structural_verification_required=True,
                visual_verification_required=False,
                render_required=False,
                postconditions=["output sheet count == input sheet count", "requested metadata fields match"],
            ),
        }

    def capabilities(self) -> list[Capability]:
        caps = []
        if _has("openpyxl"):
            import importlib.metadata as im

            try:
                version = im.version("openpyxl")
            except im.PackageNotFoundError:
                version = None
            caps.append(
                Capability(
                    id="xlsx.structural",
                    status=CapabilityStatus.AVAILABLE,
                    detail="openpyxl importable: structural inspect/verify/metadata_set available. "
                    "Formula recalculation is never available via this backend (see this module's docstring).",
                    detected_via="import openpyxl",
                    version=version,
                )
            )
        else:
            caps.append(
                Capability(id="xlsx.structural", status=CapabilityStatus.MISSING, detail="openpyxl not importable.")
            )

        soffice = soffice_binary()
        if soffice:
            caps.append(
                Capability(
                    id="xlsx.render",
                    status=CapabilityStatus.AVAILABLE,
                    detail=f"LibreOffice found on PATH ({soffice}). See adapters/pptx/adapter.py's docstring: "
                    "a present binary does not guarantee a specific document converts successfully.",
                    detected_via="which soffice",
                )
            )
        else:
            caps.append(
                Capability(
                    id="xlsx.render",
                    status=CapabilityStatus.MISSING,
                    detail="No soffice/libreoffice binary found on PATH.",
                    detected_via="which soffice",
                )
            )
        return caps

    def limitations(self) -> list[str]:
        return [
            "Formulas are never recalculated (openpyxl cannot evaluate formulas); cached results, when "
            "present, reflect whatever application last saved the file, not this adapter's own computation.",
            "Chart and embedded-drawing validity is not checked.",
            "Conditional formatting and data validation rules are not checked.",
            "Rendering depends on an external LibreOffice install; a present binary does not guarantee "
            "a specific document converts successfully.",
        ]

    def recognized_policy_keys(self) -> frozenset[str]:
        return frozenset({
            "require_sheet_count", "min_sheets", "max_sheets", "require_sheet_names", "forbid_placeholder_text",
            "forbid_external_links", "require_metadata",
        })

    # ---- inspect ---------------------------------------------------

    def inspect(self, ref: ArtifactRef) -> InspectionReport:
        check_input_size(ref.path)
        openpyxl = _require_openpyxl()
        _reject_entities_before_opening(ref.path)
        warnings: list[str] = []
        try:
            wb = openpyxl.load_workbook(str(ref.path), data_only=False)
        except Exception as exc:
            raise ArtifactInputError(
                code="ARTIFACT_XLSX_UNREADABLE",
                message=f"openpyxl could not open '{ref.path}': {exc}",
                remediation="The file may be corrupt or not a valid XLSX despite its OOXML content type.",
                evidence={"path": str(ref.path)},
            ) from exc

        try:
            wb_cached = openpyxl.load_workbook(str(ref.path), data_only=True)
        except Exception:  # noqa: BLE001 - the formula-mode load above already succeeded; treat as no cache
            wb_cached = None

        sheet_names = wb.sheetnames
        formula_cells = 0
        cached_errors: list[str] = []
        cached_values_seen = False
        all_text_parts: list[str] = []

        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            ws_cached = wb_cached[sheet_name] if wb_cached is not None else None
            for row in ws.iter_rows():
                for cell in row:
                    if cell.data_type == "f":
                        formula_cells += 1
                        if ws_cached is not None:
                            cached_value = ws_cached[cell.coordinate].value
                            if cached_value is not None:
                                cached_values_seen = True
                                if isinstance(cached_value, str) and cached_value.strip() in _ERROR_TOKENS:
                                    cached_errors.append(f"{sheet_name}!{cell.coordinate}: {cached_value}")
                    elif isinstance(cell.value, str):
                        all_text_parts.append(cell.value)

        external_links = list(getattr(wb, "_external_links", []) or [])
        defined_names = list(wb.defined_names.keys()) if hasattr(wb.defined_names, "keys") else list(wb.defined_names)

        props = wb.properties
        metadata = {
            "title": props.title or "",
            "author": props.creator or "",
            "subject": props.subject or "",
            "keywords": props.keywords or "",
        }

        details = {
            "sheet_count": len(sheet_names),
            "sheet_names": sheet_names,
            "formula_cells": formula_cells,
            "cached_values_seen": cached_values_seen,
            "cached_errors": cached_errors,
            "external_links": [str(link) for link in external_links],
            "defined_names": defined_names,
            "metadata": metadata,
            # Same rationale as the other adapters' leftover_markers: the
            # marker list found, not every cell's text content.
            "leftover_markers": find_leftover_markers("\n".join(all_text_parts)),
        }
        return InspectionReport(artifact=ref, details=details, warnings=warnings)

    # ---- plan --------------------------------------------------------

    def plan(self, ref: ArtifactRef, operation: str, args: dict[str, Any], output_path: Path) -> OperationPlan:
        specs = self.operations()
        if operation not in specs:
            raise ArtifactInputError(
                code="ARTIFACT_OPERATION_UNKNOWN",
                message=f"XLSX adapter has no operation '{operation}'.",
                remediation=f"Supported operations: {sorted(specs)}.",
                evidence={"operation": operation},
            )
        spec = specs[operation]
        report = self.inspect(ref)
        risks: list[str] = []
        if report.details["external_links"]:
            risks.append(f"{len(report.details['external_links'])} external workbook link(s) present.")
        if report.details["cached_errors"]:
            risks.append(f"{len(report.details['cached_errors'])} cell(s) have a cached formula error.")

        return OperationPlan(
            operation=f"xlsx.{operation}",
            adapter=self.id,
            input=ref.to_dict(),
            output_path=str(output_path),
            required_capabilities=["xlsx.structural"] + (["xlsx.render"] if spec.render_required else []),
            files_touched=[],
            files_created=[str(output_path)],
            rendering_strategy="soffice -> PDF -> pypdfium2 page images" if spec.render_required else None,
            verification_strategy={
                "structural_required": spec.structural_verification_required,
                "visual_required": spec.visual_verification_required,
            },
            risks=risks,
            warnings=list(report.warnings),
        )

    # ---- execute -------------------------------------------------------

    def execute(self, ref: ArtifactRef, operation: str, args: dict[str, Any], output_path: Path) -> ArtifactRef:
        openpyxl = _require_openpyxl()
        if operation != "metadata_set":
            raise ArtifactInputError(
                code="ARTIFACT_OPERATION_UNKNOWN",
                message=f"XLSX adapter has no operation '{operation}'.",
                evidence={"operation": operation},
            )
        _reject_entities_before_opening(ref.path)
        wb = openpyxl.load_workbook(str(ref.path))
        props = wb.properties
        if "title" in args:
            props.title = args["title"]
        if "author" in args:
            props.creator = args["author"]
        if "subject" in args:
            props.subject = args["subject"]
        if "keywords" in args:
            props.keywords = args["keywords"]

        output_path.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.NamedTemporaryFile(dir=str(output_path.parent), suffix=".xlsx.tmp", delete=False) as tmp:
            tmp_path = Path(tmp.name)
        try:
            wb.save(str(tmp_path))
            atomic_copy(tmp_path, output_path)
        finally:
            tmp_path.unlink(missing_ok=True)
        return ArtifactRef.from_path(output_path)

    # ---- render ----------------------------------------------------

    def render(self, ref: ArtifactRef, out_dir: Path, *, limits: Limits = DEFAULT_LIMITS) -> RenderResult:
        with tempfile.TemporaryDirectory(prefix="artifacts-skill-xlsx-render-") as tmp:
            pdf_path = convert_to_pdf(ref.path, Path(tmp) / "pdf", limits=limits)
            return render_pdf_pages(pdf_path, out_dir, limits=limits)

    # ---- verify ------------------------------------------------------

    def verify_structural(self, ref: ArtifactRef, policy: dict[str, Any]) -> VerificationResult:
        checks: list[Check] = []
        try:
            report = self.inspect(ref)
        except ArtifactInputError as exc:
            checks.append(Check(id="xlsx_validity", name="XLSX is readable", status=CheckStatus.FAIL, message=str(exc)))
            return VerificationResult(kind="structural", checks=checks)

        checks.append(Check(id="xlsx_validity", name="XLSX is readable", status=CheckStatus.PASS))

        details = report.details
        sheet_count = details["sheet_count"]

        checks.append(
            Check(
                id="sheet_count",
                name="Sheet count",
                status=CheckStatus.FAIL if sheet_count == 0 else CheckStatus.PASS,
                message=f"{sheet_count} sheet(s).",
                evidence={"sheet_count": sheet_count},
            )
        )

        if "require_sheet_count" in policy:
            expected = policy["require_sheet_count"]
            ok = sheet_count == expected
            checks.append(
                Check(
                    id="sheet_count_requirement",
                    name="Sheet count matches requirement",
                    status=CheckStatus.PASS if ok else CheckStatus.FAIL,
                    message=f"expected {expected}, got {sheet_count}.",
                )
            )

        if "min_sheets" in policy or "max_sheets" in policy:
            lo = policy.get("min_sheets", 0)
            hi = policy.get("max_sheets", float("inf"))
            ok = lo <= sheet_count <= hi
            checks.append(
                Check(
                    id="sheet_count_range",
                    name="Sheet count within range",
                    status=CheckStatus.PASS if ok else CheckStatus.FAIL,
                    message=f"expected [{lo}, {hi}], got {sheet_count}.",
                )
            )

        if "require_sheet_names" in policy:
            expected_names = set(policy["require_sheet_names"])
            actual_names = set(details["sheet_names"])
            ok = expected_names.issubset(actual_names)
            checks.append(
                Check(
                    id="required_sheet_names",
                    name="Required sheet names present",
                    status=CheckStatus.PASS if ok else CheckStatus.FAIL,
                    message=f"missing: {sorted(expected_names - actual_names)}" if not ok else "all present.",
                )
            )

        leftover_markers = details.get("leftover_markers", [])
        if leftover_markers:
            checks.append(
                Check(
                    id="leftover_placeholder_text",
                    name="No leftover generation placeholder text",
                    status=CheckStatus.FAIL if policy.get("forbid_placeholder_text") else CheckStatus.WARN,
                    message=f"Found likely-unreviewed placeholder text: {leftover_markers}.",
                    evidence={"markers": leftover_markers},
                )
            )
        else:
            checks.append(
                Check(id="leftover_placeholder_text", name="No leftover generation placeholder text", status=CheckStatus.PASS)
            )

        if details["external_links"]:
            checks.append(
                Check(
                    id="external_links",
                    name="No external workbook links",
                    status=CheckStatus.FAIL if policy.get("forbid_external_links") else CheckStatus.WARN,
                    message=f"{len(details['external_links'])} external link(s) present; not fetched (network "
                    "access is off by default — see docs/security.md).",
                    evidence={"links": details["external_links"]},
                )
            )
        else:
            checks.append(Check(id="external_links", name="No external workbook links", status=CheckStatus.PASS))

        if details["formula_cells"] == 0:
            checks.append(
                Check(
                    id="formula_cached_errors",
                    name="No cached formula errors",
                    status=CheckStatus.SKIPPED,
                    message="No formulas present.",
                )
            )
            checks.append(
                Check(
                    id="formula_recalculation",
                    name="Formula recalculation",
                    status=CheckStatus.SKIPPED,
                    message="No formulas present.",
                )
            )
        else:
            if details["cached_errors"]:
                checks.append(
                    Check(
                        id="formula_cached_errors",
                        name="No cached formula errors",
                        status=CheckStatus.FAIL,
                        message=f"{len(details['cached_errors'])} cell(s) have a cached error value.",
                        evidence={"cells": details["cached_errors"]},
                    )
                )
            elif details["cached_values_seen"]:
                checks.append(
                    Check(
                        id="formula_cached_errors",
                        name="No cached formula errors",
                        status=CheckStatus.PASS,
                        message="Cached results present and none are error values.",
                    )
                )
            else:
                checks.append(
                    Check(
                        id="formula_cached_errors",
                        name="No cached formula errors",
                        status=CheckStatus.UNKNOWN,
                        message="Formulas are present but no cached results exist to inspect (the file was "
                        "likely written without ever being opened by a spreadsheet application).",
                    )
                )
            checks.append(
                Check(
                    id="formula_recalculation",
                    name="Formula recalculation",
                    status=CheckStatus.UNKNOWN,
                    message="openpyxl does not evaluate formulas; freshness of any cached results relative "
                    "to current cell inputs cannot be determined by this adapter (see this module's docstring).",
                )
            )

        if "require_metadata" in policy:
            for key, expected_value in policy["require_metadata"].items():
                actual = details["metadata"].get(key.lower())
                ok = actual == expected_value
                checks.append(
                    Check(
                        id=f"metadata_{key.lower()}",
                        name=f"Metadata field '{key}'",
                        status=CheckStatus.PASS if ok else CheckStatus.FAIL,
                        message=f"expected '{expected_value}', got '{actual}'.",
                    )
                )

        return VerificationResult(kind="structural", checks=checks)
